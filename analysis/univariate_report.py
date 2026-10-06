"""
Phase 2 report: a single factor analysis (SFA) workbook, one sheet per feature.

Reads the three CSVs written by analysis/univariate.py and builds
outputs/univariate_report.xlsx with native Excel charts:
  - Summary sheet: every feature sorted by IV, flags colour-coded, links to each sheet
  - One sheet per feature:
      * bin table: cars per bin, share, fail rate, WoE (train) and validation fail rate
      * chart 1: cars per bin (bars) with train and validation fail rate (lines)
      * chart 2: weight of evidence per bin
      * chart 3: fail rate per bin by test year (is the pattern stable over time?)

Run from the repository root, after analysis.univariate:
    python -m analysis.univariate_report
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import OUTPUTS

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
RED = PatternFill("solid", fgColor="F8CBAD")
AMBER = PatternFill("solid", fgColor="FFE699")
GREEN = PatternFill("solid", fgColor="C6EFCE")
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)


def style_header(row):
    for c in row:
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BORDER


def style_body(ws, min_row, max_row, max_col):
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, max_col=max_col):
        for c in row:
            c.font = Font(name=FONT, size=10)
            c.border = BORDER


def flag_fill(flags: str):
    if any(k in flags for k in ("IV<0.02", "IV>0.5", "PSI>0.25")):
        return RED
    if any(k in flags for k in ("monitor", "non-monotonic")):
        return AMBER
    return GREEN


def one_page(ws, fit_height=True):
    """Print the sheet on one landscape page (handy for a PDF of the whole pack)."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1 if fit_height else 0


def show_axes(chart):
    # Newer openpyxl versions hide axes unless told otherwise
    chart.x_axis.delete = False
    chart.y_axis.delete = False


def summary_sheet(wb, summary: pd.DataFrame):
    ws = wb.active
    ws.title = "Summary"
    cols = ["feature", "IV", "AUC_train", "AUC_val", "PSI", "missing_pct",
            "distinct", "bins", "monotonic", "flags"]
    headers = ["Feature", "IV", "AUC (train)", "AUC (validation)", "PSI", "Missing %",
               "Distinct values", "Bins", "Monotonicity", "Flags"]
    ws.append(headers)
    style_header(ws[1])
    for _, r in summary.iterrows():
        ws.append([r[c] if not (isinstance(r[c], float) and pd.isna(r[c])) else "" for c in cols])
    style_body(ws, 2, ws.max_row, len(cols))
    for row in range(2, ws.max_row + 1):
        name = ws.cell(row=row, column=1)
        name.hyperlink = f"#'{name.value}'!A1"
        name.font = Font(name=FONT, size=10, color="1F4E79", underline="single")
        flags = str(ws.cell(row=row, column=10).value or "")
        ws.cell(row=row, column=10).fill = flag_fill(flags)
        ws.cell(row=row, column=2).number_format = "0.000"
        ws.cell(row=row, column=5).number_format = "0.000"
    for i, w in enumerate([28, 9, 11, 13, 9, 10, 10, 7, 16, 48], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    one_page(ws, fit_height=False)
    n = ws.max_row + 2
    notes = [
        "IV: <0.02 useless, 0.02-0.1 weak, 0.1-0.3 medium, >0.3 strong, >0.5 check for leakage",
        "PSI (train vs validation): <0.1 stable, 0.1-0.25 monitor, >0.25 unstable",
        "Monotonicity: Spearman correlation of fail rate with bin order, bins of 100+ cars only",
        "Flags colour: red = fails a gate, amber = investigate, green = passes",
        "Click a feature name to open its sheet.",
    ]
    for k, text in enumerate(notes):
        ws.cell(row=n + k, column=1, value=text).font = Font(name=FONT, size=9, italic=True)


def feature_sheet(wb, feat: str, s: pd.Series, bins: pd.DataFrame, by_year: pd.DataFrame):
    ws = wb.create_sheet(title=feat[:31])
    ws["A1"] = feat
    ws["A1"].font = Font(name=FONT, size=14, bold=True)
    ws["A2"] = (f"IV {s['IV']:.3f}  |  AUC train {s['AUC_train']:.3f}, validation {s['AUC_val']:.3f}  |  "
                f"PSI {s['PSI']:.3f}  |  {s['monotonic']}  |  missing {s['missing_pct']}%")
    ws["A2"].font = Font(name=FONT, size=10)
    flags = s["flags"] if isinstance(s["flags"], str) else ""
    ws["A3"] = flags or "No flags"
    ws["A3"].font = Font(name=FONT, size=10, bold=True)
    ws["A3"].fill = flag_fill(flags)
    ws["H1"] = "Back to summary"
    ws["H1"].hyperlink = "#'Summary'!A1"
    ws["H1"].font = Font(name=FONT, size=10, color="1F4E79", underline="single")

    # ---- Bin table (row 5 onwards) ----
    top = 5
    headers = ["Bin", "Cars (train)", "Share", "Fail rate (train)", "WoE",
               "Cars (validation)", "Fail rate (validation)"]
    for j, h in enumerate(headers, 1):
        ws.cell(row=top, column=j, value=h)
    style_header(ws[top])
    total = bins["n_train"].sum()
    for k, (_, b) in enumerate(bins.iterrows(), 1):
        r = top + k
        ws.cell(row=r, column=1, value=str(b["bin"]))
        ws.cell(row=r, column=2, value=int(b["n_train"]))
        ws.cell(row=r, column=3, value=b["n_train"] / total).number_format = "0.0%"
        ws.cell(row=r, column=4, value=b["fail_rate_train"]).number_format = "0.0%"
        ws.cell(row=r, column=5, value=b["woe"]).number_format = "0.000"
        if pd.notna(b["n_val"]):
            ws.cell(row=r, column=6, value=int(b["n_val"]))
        if pd.notna(b["fail_rate_val"]):
            ws.cell(row=r, column=7, value=b["fail_rate_val"]).number_format = "0.0%"
    last = top + len(bins)
    style_body(ws, top + 1, last, len(headers))
    for i, w in enumerate([22, 12, 9, 14, 9, 14, 16], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    cats = Reference(ws, min_col=1, min_row=top + 1, max_row=last)

    # ---- Chart 1: cars per bin + fail rates ----
    bar = BarChart()
    bar.title = "Cars per bin and fail rate"
    bar.add_data(Reference(ws, min_col=2, min_row=top, max_row=last), titles_from_data=True)
    bar.set_categories(cats)
    bar.y_axis.title = "Cars (train)"
    bar.y_axis.majorGridlines = None
    line = LineChart()
    line.add_data(Reference(ws, min_col=4, min_row=top, max_row=last), titles_from_data=True)
    line.add_data(Reference(ws, min_col=7, min_row=top, max_row=last), titles_from_data=True)
    line.y_axis.axId = 200
    line.y_axis.title = "Fail rate"
    line.y_axis.number_format = "0%"
    line.y_axis.crosses = "max"
    for series in line.series:
        series.smooth = False          # straight lines between bins: no invented curves
    show_axes(bar)
    show_axes(line)
    bar += line
    bar.height, bar.width = 8, 16
    ws.add_chart(bar, "I5")

    # ---- Chart 2: WoE per bin ----
    woe = BarChart()
    woe.title = "Weight of evidence (negative = riskier)"
    woe.add_data(Reference(ws, min_col=5, min_row=top, max_row=last), titles_from_data=True)
    woe.set_categories(cats)
    woe.legend = None
    woe.y_axis.title = "WoE"
    woe.x_axis.tickLblPos = "low"      # bin labels below the chart, not on the bars
    show_axes(woe)
    woe.height, woe.width = 8, 16
    ws.add_chart(woe, "I22")

    # ---- Chart 3: fail rate per bin by year ----
    pivot = by_year.pivot_table(index="bin", columns="year", values="fail_rate")
    pivot = pivot.reindex([b for b in bins["bin"].astype(str) if b in pivot.index])
    yr_top = last + 3
    ws.cell(row=yr_top - 1, column=1, value="Fail rate by test year (train period)").font = \
        Font(name=FONT, size=11, bold=True)
    ws.cell(row=yr_top, column=1, value="Year")
    for j, b in enumerate(pivot.index, 2):
        ws.cell(row=yr_top, column=j, value=str(b))
    style_header(ws[yr_top])
    for k, year in enumerate(pivot.columns, 1):
        ws.cell(row=yr_top + k, column=1, value=str(year))
        for j, b in enumerate(pivot.index, 2):
            v = pivot.loc[b, year]
            if pd.notna(v):
                ws.cell(row=yr_top + k, column=j, value=float(v)).number_format = "0.0%"
    yr_last = yr_top + len(pivot.columns)
    style_body(ws, yr_top + 1, yr_last, len(pivot.index) + 1)
    if len(pivot.index) and len(pivot.columns) > 1:
        yl = LineChart()
        yl.title = "Fail rate per bin by year"
        yl.add_data(Reference(ws, min_col=2, max_col=len(pivot.index) + 1,
                              min_row=yr_top, max_row=yr_last), titles_from_data=True)
        yl.set_categories(Reference(ws, min_col=1, min_row=yr_top + 1, max_row=yr_last))
        yl.y_axis.number_format = "0%"
        yl.y_axis.title = "Fail rate"
        for series in yl.series:
            series.smooth = False
        show_axes(yl)
        yl.height, yl.width = 8, 16
        ws.add_chart(yl, "I39")
    one_page(ws)


def main() -> None:
    files = {k: OUTPUTS / f"univariate_{k}.csv" for k in ("summary", "bins", "by_year")}
    missing = [str(p) for p in files.values() if not p.exists()]
    if missing:
        raise SystemExit(f"Run python -m analysis.univariate first. Missing: {missing}")
    summary = pd.read_csv(files["summary"])
    bins = pd.read_csv(files["bins"], dtype={"bin": str})
    by_year = pd.read_csv(files["by_year"], dtype={"bin": str, "year": str})

    wb = Workbook()
    summary_sheet(wb, summary)
    for _, s in summary.iterrows():
        feat = s["feature"]
        feature_sheet(wb, feat, s,
                      bins[bins["feature"] == feat].reset_index(drop=True),
                      by_year[by_year["feature"] == feat])
    out = OUTPUTS / "univariate_report.xlsx"
    wb.save(out)
    print(f"Saved {out} ({len(summary)} feature sheets)")


if __name__ == "__main__":
    main()
